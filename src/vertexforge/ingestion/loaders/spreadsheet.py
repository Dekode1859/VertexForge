"""Deterministic spreadsheet ingestion with trusted table artifacts."""

from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from vertexforge.ingestion.loaders.base import BaseLoader
from vertexforge.ingestion.models import ArtifactRecord, IngestionResult


def _normalize_cell(value: Any) -> Any:
    """Normalize spreadsheet cell values for JSON serialization."""
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    return str(value)


def _rows_to_markdown(rows: list[list[Any]]) -> str:
    """Build a simple markdown table preview."""
    nonempty_rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    if not nonempty_rows:
        return ""

    header = ["" if cell is None else str(cell) for cell in nonempty_rows[0]]
    separator = ["---"] * len(header)
    body = [
        [" " if cell is None else str(cell) for cell in row]
        for row in nonempty_rows[1: min(len(nonempty_rows), 6)]
    ]

    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _rows_to_csv(rows: list[list[Any]]) -> str:
    """Serialize rows to CSV text."""
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerows(rows)
    return buffer.getvalue()


class SpreadsheetLoader(BaseLoader):
    """Loader for CSV, TSV, and XLSX spreadsheet files."""

    name = "spreadsheet_loader"
    detected_type = "spreadsheet"
    supported_extensions = (".csv", ".tsv", ".xlsx")

    def load(self, path: Path) -> IngestionResult:
        """Read the spreadsheet and emit sheet and table artifacts."""
        source = self.build_source(path)
        if path.suffix.lower() == ".xlsx":
            sheet_payloads = self._load_xlsx(path)
        else:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
            sheet_payloads = [self._load_delimited(path, delimiter)]

        artifacts: list[ArtifactRecord] = []
        spreadsheet_artifact = ArtifactRecord.create(
            source_id=source.source_id,
            artifact_type="document",
            subtype="spreadsheet_workbook",
            title=source.display_name,
            preview_text=f"{len(sheet_payloads)} sheet(s)",
            payload={
                "sheet_count": len(sheet_payloads),
                "contains_tables": any(sheet["row_count"] > 0 for sheet in sheet_payloads),
                "contains_extractable_text": any(sheet["row_count"] > 0 for sheet in sheet_payloads),
            },
            provenance={
                "source_filename": source.filename,
                "origin_loader": self.name,
                "extraction_method": "structured_sheet_parse",
            },
            extraction={"confidence": 1.0, "warnings": [], "fallback_used": False},
        )
        artifacts.append(spreadsheet_artifact)

        for index, sheet in enumerate(sheet_payloads):
            sheet_artifact = ArtifactRecord.create(
                source_id=source.source_id,
                parent_artifact_id=spreadsheet_artifact.artifact_id,
                artifact_type="sheet",
                subtype="spreadsheet_sheet",
                title=sheet["sheet_name"],
                preview_text=sheet["markdown"][:200] or None,
                payload={
                    "sheet_name": sheet["sheet_name"],
                    "row_count": sheet["row_count"],
                    "column_count": sheet["column_count"],
                },
                provenance={
                    "source_filename": source.filename,
                    "sheet_name": sheet["sheet_name"],
                    "sheet_index": index,
                    "origin_loader": self.name,
                    "extraction_method": "structured_sheet_parse",
                },
                extraction={"confidence": 1.0, "warnings": [], "fallback_used": False},
            )
            table_artifact = ArtifactRecord.create(
                source_id=source.source_id,
                parent_artifact_id=sheet_artifact.artifact_id,
                artifact_type="table",
                subtype="spreadsheet_sheet_table",
                title=sheet["sheet_name"],
                preview_text=sheet["markdown"][:200] or None,
                payload={
                    "sheet_name": sheet["sheet_name"],
                    "headers": sheet["headers"],
                    "rows": sheet["rows"],
                    "markdown": sheet["markdown"],
                    "csv": sheet["csv"],
                },
                provenance={
                    "source_filename": source.filename,
                    "sheet_name": sheet["sheet_name"],
                    "origin_loader": self.name,
                    "extraction_method": "structured_sheet_parse",
                },
                extraction={"confidence": 1.0, "warnings": [], "fallback_used": False},
            )
            artifacts.extend([sheet_artifact, table_artifact])

        return IngestionResult(source=source, artifacts=artifacts)

    def _load_delimited(self, path: Path, delimiter: str) -> dict[str, Any]:
        """Load CSV or TSV files."""
        with path.open("r", encoding="utf-8", newline="") as handle:
            rows = [[_normalize_cell(cell) for cell in row] for row in csv.reader(handle, delimiter=delimiter)]
        return self._build_sheet_payload(path.stem, rows)

    def _load_xlsx(self, path: Path) -> list[dict[str, Any]]:
        """Load rows from each workbook sheet."""
        workbook = load_workbook(filename=path, data_only=True)
        sheets: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            rows = [[_normalize_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            sheets.append(self._build_sheet_payload(sheet.title, rows))
        return sheets

    def _build_sheet_payload(self, sheet_name: str, rows: list[list[Any]]) -> dict[str, Any]:
        """Construct normalized sheet data."""
        nonempty_rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
        headers = []
        if nonempty_rows:
            headers = ["" if cell is None else str(cell) for cell in nonempty_rows[0]]

        return {
            "sheet_name": sheet_name,
            "row_count": len(nonempty_rows),
            "column_count": max((len(row) for row in nonempty_rows), default=0),
            "headers": headers,
            "rows": nonempty_rows[1:] if len(nonempty_rows) > 1 else [],
            "markdown": _rows_to_markdown(nonempty_rows),
            "csv": _rows_to_csv(nonempty_rows),
        }

